"""
Word Document Exporter - Export personalized resumes to .docx format
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from .models import PersonalizedResume, Resume
from typing import Optional


class WordDocumentExporter:
    """Export resume to Word (.docx) format"""
    
    def __init__(self):
        self.doc = None
    
    def export_resume(self, 
                     personalized_resume: PersonalizedResume,
                     output_path: str,
                     include_matching_score: bool = True) -> str:
        """
        Export personalized resume to Word document
        
        Args:
            personalized_resume: PersonalizedResume object
            output_path: Full path where to save the .docx file
            include_matching_score: Whether to include job matching score
            
        Returns:
            Path to the created file
        """
        self.doc = Document()
        
        # Set up document margins
        sections = self.doc.sections
        for section in sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)
        
        resume = personalized_resume.original_resume
        
        # Add header with name and title
        self._add_header(resume, personalized_resume)
        
        # Add matching info if requested
        if include_matching_score:
            self._add_matching_info(personalized_resume)
        
        # Add contact information
        self._add_contact_info(resume)
        
        # Add professional summary
        self._add_professional_summary(personalized_resume)
        
        # Add skills section
        self._add_skills_section(personalized_resume)
        
        # Add work experience
        self._add_experience_section(personalized_resume)
        
        # Add education
        self._add_education_section(resume)
        
        # Add certifications if available
        if resume.certifications:
            self._add_certifications_section(resume)
        
        # Save document
        self.doc.save(output_path)
        return output_path
    
    def _add_header(self, resume: Resume, personalized_resume: PersonalizedResume):
        """Add resume header with name and job title"""
        # Name
        name_para = self.doc.add_paragraph()
        name_run = name_para.add_run(resume.name)
        name_run.font.size = Pt(16)
        name_run.font.bold = True
        name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Job title and company
        job_para = self.doc.add_paragraph()
        job_run = job_para.add_run(
            f"{personalized_resume.target_job.job_title} @ {personalized_resume.target_job.company}"
        )
        job_run.font.size = Pt(11)
        job_run.font.italic = True
        job_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        self.doc.add_paragraph()  # Spacing
    
    def _add_matching_info(self, personalized_resume: PersonalizedResume):
        """Add job matching score and information"""
        info_para = self.doc.add_paragraph()
        score_run = info_para.add_run(
            f"Job Match Score: {personalized_resume.relevance_score:.1f}%"
        )
        score_run.font.bold = True
        score_run.font.color.rgb = RGBColor(0, 102, 0)
        
        self.doc.add_paragraph()  # Spacing
    
    def _add_contact_info(self, resume: Resume):
        """Add contact information section"""
        contact_para = self.doc.add_paragraph()
        contact_run = contact_para.add_run(resume.email)
        contact_run.font.size = Pt(10)
        
        contact_para.add_run(f" | {resume.phone}").font.size = Pt(10)
        contact_para.add_run(f" | {resume.location}").font.size = Pt(10)
        
        self.doc.add_paragraph()  # Spacing
    
    def _add_professional_summary(self, personalized_resume: PersonalizedResume):
        """Add professional summary section"""
        if not personalized_resume.personalized_summary:
            return
        
        heading = self.doc.add_heading('PROFESSIONAL SUMMARY', level=2)
        heading.runs[0].font.size = Pt(11)
        heading.runs[0].font.bold = True
        
        summary_para = self.doc.add_paragraph(personalized_resume.personalized_summary)
        summary_para.paragraph_format.space_after = Pt(6)
        
        self.doc.add_paragraph()  # Spacing
    
    def _add_skills_section(self, personalized_resume: PersonalizedResume):
        """Add technical skills section"""
        if not personalized_resume.matched_skills:
            return
        
        heading = self.doc.add_heading('TECHNICAL SKILLS', level=2)
        heading.runs[0].font.size = Pt(11)
        heading.runs[0].font.bold = True
        
        # Group skills by category if available
        skills_by_category = {}
        for skill in personalized_resume.matched_skills:
            category = skill.category or 'Technical'
            if category not in skills_by_category:
                skills_by_category[category] = []
            skills_by_category[category].append(skill)
        
        for category, skills in skills_by_category.items():
            category_para = self.doc.add_paragraph()
            category_run = category_para.add_run(f"{category}: ")
            category_run.bold = True
            
            skill_names = [f"{s.name} ({s.proficiency_level})" for s in skills]
            category_para.add_run(', '.join(skill_names))
        
        self.doc.add_paragraph()  # Spacing
    
    def _add_experience_section(self, personalized_resume: PersonalizedResume):
        """Add work experience section"""
        if not personalized_resume.prioritized_experience:
            return
        
        heading = self.doc.add_heading('PROFESSIONAL EXPERIENCE', level=2)
        heading.runs[0].font.size = Pt(11)
        heading.runs[0].font.bold = True
        
        for exp in personalized_resume.prioritized_experience:
            # Job title and company
            title_para = self.doc.add_paragraph()
            title_run = title_para.add_run(exp.job_title)
            title_run.bold = True
            title_run.font.size = Pt(10)
            
            # Company and dates
            company_para = self.doc.add_paragraph(
                f"{exp.company} | {exp.start_date} - {exp.end_date or 'Present'}"
            )
            company_para.paragraph_format.left_indent = Inches(0.25)
            company_para.runs[0].font.italic = True
            company_para.runs[0].font.size = Pt(9)
            
            # Description
            if exp.description:
                desc_para = self.doc.add_paragraph(exp.description)
                desc_para.paragraph_format.left_indent = Inches(0.25)
                desc_para.paragraph_format.space_after = Pt(3)
            
            # Achievements as bullet points
            if exp.achievements:
                for achievement in exp.achievements:
                    achievement_para = self.doc.add_paragraph(
                        achievement,
                        style='List Bullet'
                    )
                    achievement_para.paragraph_format.left_indent = Inches(0.5)
                    achievement_para.paragraph_format.space_after = Pt(2)
            
            self.doc.add_paragraph()  # Spacing between jobs
    
    def _add_education_section(self, resume: Resume):
        """Add education section"""
        if not resume.education:
            return
        
        heading = self.doc.add_heading('EDUCATION', level=2)
        heading.runs[0].font.size = Pt(11)
        heading.runs[0].font.bold = True
        
        for edu in resume.education:
            # Degree
            degree_para = self.doc.add_paragraph()
            degree_run = degree_para.add_run(f"{edu.degree} in {edu.field_of_study}")
            degree_run.bold = True
            degree_run.font.size = Pt(10)
            
            # Institution and graduation year
            inst_para = self.doc.add_paragraph(f"{edu.institution}, {edu.graduation_year}")
            inst_para.paragraph_format.left_indent = Inches(0.25)
            inst_para.runs[0].font.italic = True
            inst_para.runs[0].font.size = Pt(9)
            
            if edu.relevant_coursework:
                coursework_para = self.doc.add_paragraph()
                coursework_para.add_run("Relevant Coursework: ").bold = True
                coursework_para.add_run(', '.join(edu.relevant_coursework))
                coursework_para.paragraph_format.left_indent = Inches(0.25)
            
            self.doc.add_paragraph()  # Spacing
    
    def _add_certifications_section(self, resume: Resume):
        """Add certifications section"""
        if not resume.certifications:
            return
        
        heading = self.doc.add_heading('CERTIFICATIONS', level=2)
        heading.runs[0].font.size = Pt(11)
        heading.runs[0].font.bold = True
        
        for cert in resume.certifications:
            cert_para = self.doc.add_paragraph(cert, style='List Bullet')
            cert_para.runs[0].font.size = Pt(10)
        
        self.doc.add_paragraph()  # Spacing


def create_tailored_resume_word(
    personalized_resume: PersonalizedResume,
    output_path: str,
    include_matching_score: bool = True
) -> str:
    """
    Convenience function to create a tailored resume in Word format
    
    Args:
        personalized_resume: PersonalizedResume object
        output_path: Full path where to save the .docx file
        include_matching_score: Whether to include job matching score
        
    Returns:
        Path to the created file
    """
    exporter = WordDocumentExporter()
    return exporter.export_resume(personalized_resume, output_path, include_matching_score)
